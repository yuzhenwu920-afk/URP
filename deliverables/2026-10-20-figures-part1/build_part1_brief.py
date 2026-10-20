"""Build the reviewable bilingual figure-editing brief and its archive."""
from pathlib import Path
import json,hashlib,zipfile
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
F1_CAPTION=('Figure 1. Workflow linking two-stage MCMC calibration, global sensitivity analysis, and irrigation optimization for maize during 2019–2022. '
    'Sobol analysis uses eight independently distributed parameters within common uniform ranges, whereas NSGA-III searches for irrigation schedules at frozen year-specific parameter centers. '
    'NSGA-II provides a same-budget benchmark. Candidate schedules are assessed using stage-specific conditional perturbations and seasonal water balances, including terminal soil water storage. '
    'Dashed arrows indicate the transfer of frozen parameter centers or stage-specific samples; solid arrows indicate model or data flow. '
    'Stage-specific inference and sampling-diagnostic limitations are retained.')
F10_CAPTION=('Figure 10. Daily reference evapotranspiration (ET₀) and precipitation during the simulation periods in (a) 2019, (b) 2020, (c) 2021, and (d) 2022. '
    'Colored lines show ET₀ (left axis), and gray bars show daily precipitation totals (right axis). '
    'Each panel covers the original calendar window used by the corresponding annual model. '
    'The ET₀ series are taken from the audited meteorological inputs used in the simulations.')
F10_LITERAL=('Colored lines show the audited reference evapotranspiration (ET₀, input set B) used by the original annual models, and gray bars show precipitation. '
    'All four years use the daily forcing records supplied in the completed data packages; annual scaling factors were not used in place of the original ET₀ series. '
    'Each panel shows the original calendar window used in the corresponding annual simulation.')
BODY=('Daily reference evapotranspiration and precipitation used to drive the annual crop models are presented in Figure 10. '
    'The records retain their original daily resolution and year-specific simulation windows, providing the meteorological context for comparisons among years. '
    'Reference evapotranspiration characterizes atmospheric evaporative demand; actual crop water use and soil water availability are evaluated separately from the simulated water balances.')
TRANSLATIONS=[('四年实际使用的气象驱动','Meteorological forcing used in the four annual simulations'),
    ('环境条件（本图建议具体写成气象条件）','Meteorological conditions during the simulation periods (2019–2022)'),
    ('参考蒸散 ETo（mm/d）','Reference evapotranspiration, ET₀ (mm d⁻¹)'),
    ('降水（mm）','Precipitation (mm)'),('日期','Date (MM–DD)'),
    ('参考蒸散（图例）','Reference evapotranspiration (ET₀)'),('降水（图例）','Precipitation'),
    ('2019、2020、2021、2022子图','(a) 2019; (b) 2020; (c) 2021; (d) 2022')]


def build():
    notes='''# 第一部分：Figure 1重构与Figure 10英文化

本次提供可评审图件和文字。推荐Figure1_three_column_EN作为论文双栏通栏图；Figure1_tree_EN为备选。三种图均提供600dpi PNG、矢量PDF和保留文本的SVG。

## Figure 1：布局与逻辑

推荐三列：①数据输入；②模型与分阶段校准；③敏感性、灌溉优化与候选核验。左列依次列气象、土壤与管理、校准目标；中列区分年度原模型、第一阶段五参数和第二阶段三参数；右列将Sobol和NSGA-III作为不同分析支路，优化后再做候选核验。

Sobol从年度模型分支，输入为声明的公共独立均匀范围。它不接收相关MCMC样本作为经典Sobol输入，也没有通过筛选后删掉优化模型中的参数。优化接收两阶段组合的冻结中心；核验按阶段分别接收整行冻结样本，不构造八维联合后验。

实线表示模型/数据流程，虚线表示冻结中心或分阶段样本的传递。第一阶段目标为CC、Tr、Es；第二阶段目标为生物量与产量。Tr为来源约定的目标量，不能把图中所有目标统称为独立直接实测。正文另需保留采样诊断和数字化观测误差的限制。

树状备选将数据与原模型放在顶部，下分Sobol解释支路和MCMC→优化支路，底部汇入综合解释。它更突出两条分析的独立输入定义；三列方案更容易展示数据结构及两个MCMC阶段。两种方案择一用于正文。

## Pattern的用法

- 三类模块分别用浅斜线、稀疏点纹和稀疏交叉线识别；只铺列标题左侧窄色带，正文框内保持清洁。
- 底色接近白色，纹理线使用浅灰蓝，避免黑色密纹、渐变和装饰阴影。纹理不代表误差或置信程度。
- 年份颜色继续固定为2019蓝、2020橙、2021绿、2022紫；流程阶段采用浅中性色，避免和年份编码混用。
- 以最终印刷宽度约183mm评估字号：流程图正文约8.1–9.5pt，气象图刻度9pt、轴标签约10pt、面板标题11pt。目标期刊若要求更大字号，应增大版面或缩短框内文字。
- 图内只留变量和方法关系；Sobol N512×3种子、六输出，以及优化100个体×100代×5种子的完整设置放入方法正文。

## Figure 10：英文化及数据口径

当前图展示逐日ET₀与降水，图名采用Meteorological conditions比泛指Environmental conditions更具体。ET₀译为reference evapotranspiration，不能译为actual crop evapotranspiration。降水按每日水深表示为mm，ET₀为mm d⁻¹。日期标签使用MM–DD，各年份均采用原模拟日历，未截成相同长度。

四年均采用原逐日ET₀和降水数据；不平滑、不年比例缩放，不用灌水量替代降水。左轴统一0–11mm d⁻¹，右轴统一0–60mm。相同轴范围有助于跨年比较；两轴之间的线、柱高度不作同单位比较。图注说明线与柱及各轴含义。

Figure10_daily_environment.csv导出全部599行图源；Figure10_window_summary.csv提供分年度日历、天数、ET₀合计/日均与降水合计。年窗口长度不同，合计差异应结合天数解释；当前图不据此单独推断作物干旱胁迫或温度差异。

本轮不需要新增数据。如后续把“环境条件”扩展为包含温度、土壤水或胁迫过程，应另列对应时序与测量/模拟来源；不能只根据ET₀曲线推出这些结论。

## Figure 10：中英标签对照

'''
    notes+='| 原中文 | 推荐英文 |\n|---|---|\n'+''.join(f'| {a} | {b} |\n' for a,b in TRANSLATIONS)
    notes+='\n## Figure 10：推荐投稿图注\n\n'+F10_CAPTION+'\n\n## 原图注忠实译文（供溯源）\n\n'+F10_LITERAL
    notes+='\n\n原图注关于“补齐包、年度比例”的操作记录放入方法或补充材料；投稿图注以变量、面板、单位和来源为主。\n\n## Figure 10：正文衔接示例\n\n'+BODY
    notes+='\n\n## Figure 1：英文图注\n\n'+F1_CAPTION+'\n'
    (OUT/'Figure1_Figure10_revision_notes.md').write_text(notes,encoding='utf-8')
    (OUT/'ENGLISH_CAPTIONS.txt').write_text(F1_CAPTION+'\n\n'+F10_CAPTION+'\n\nSuggested manuscript text:\n'+BODY+'\n',encoding='utf-8')
    d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(1.9);s.top_margin=s.bottom_margin=Cm(1.7)
    st=d.styles['Normal'];st.font.name='DejaVu Sans';st.font.size=Pt(10.5)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    st.paragraph_format.space_after=Pt(6)
    d.add_heading('URP：Figure 1 与 Figure 10 修订说明',0)
    d.add_paragraph('第一部分 · 2026年10月20日 · 四年数据结构、分析流程与气象条件')
    d.add_heading('1. Figure 1重构建议',1)
    for text in [
        '推荐三列结构：数据输入 → 模型与分阶段校准 → 分析与候选核验。第三列保留Sobol和优化两条不同输入的支路；两个阶段的冻结中心用于优化，分阶段样本用于条件扰动核验。',
        '树状布局作为备选：原模型之下分出Sobol支路和MCMC—优化支路，最后汇入综合解释。两种布局择一用于正文。',
        'Pattern只用于标题或模块边缘窄色带：浅斜线、稀疏点纹、稀疏交叉线。正文区域留白，阶段色使用浅中性色，保留蓝/橙/绿/紫作为四年的固定颜色。',
        '以183mm通栏宽度审阅：流程图正文约8.1–9.5pt；图内保留关系和变量，完整采样与优化预算放入方法。Sobol与优化之间不画“筛选后降维”的箭头。']:
        d.add_paragraph(text)
    d.add_heading('2. Figure 10英文标签',1)
    t=d.add_table(rows=1,cols=2);t.style='Table Grid'
    t.rows[0].cells[0].text='中文';t.rows[0].cells[1].text='推荐英文'
    for a,b in TRANSLATIONS[:5]:
        cells=t.add_row().cells;cells[0].text=a;cells[1].text=b
    d.add_heading('3. 推荐英文图注',1);d.add_paragraph(F10_CAPTION)
    d.add_heading('4. 正文衔接示例',1);d.add_paragraph(BODY)
    d.add_heading('5. 数据处理说明',1)
    d.add_paragraph('Figure 10直接保留原日历内的四年逐日ET₀和降水，共599行。未平滑、缩放或替换序列。跨年坐标范围一致。ET₀表征大气蒸发需求，实际耗水和胁迫需结合模型水量平衡解释。当前任务无需补充数据。')
    for image,title,caption in [
        ('Figure1_three_column_EN.png','Figure 1：三列推荐方案',F1_CAPTION),
        ('Figure1_tree_EN.png','Figure 1：树状备选方案','Alternative layout using two analysis branches. Use one Figure 1 layout in the manuscript.'),
        ('Figure10_environment_EN.png','Figure 10：英文版',F10_CAPTION)]:
        d.add_page_break();d.add_heading(title,1);d.add_picture(str(OUT/image),width=Cm(17.1));d.add_paragraph(caption)
    d.save(OUT/'Figure1_Figure10_revision_brief.docx')
    (OUT/'README.md').write_text('''# Figure 1 / Figure 10 第一部分修订

推荐：[Figure 1三列方案](Figure1_three_column_EN.png)；备选：[Figure 1树状方案](Figure1_tree_EN.png)。

[Figure 10英文版](Figure10_environment_EN.png) · [Word修订说明](Figure1_Figure10_revision_brief.docx) · [完整文字建议与中英对照](Figure1_Figure10_revision_notes.md) · [可复制英文图注](ENGLISH_CAPTIONS.txt)。

三幅图均有600dpi PNG、矢量PDF、可编辑文本SVG。英文环境图使用四年原逐日序列，源表及日历合计一并提供。图1提供两种备选，应择一使用。运行rebuild_figures_part1.py可重绘，已有CSV时不需要重新加载年度模型；原模型与MCMC不重跑。

[下载本部分压缩包](URP_Figures_Part1_20261020.zip)。
''',encoding='utf-8')
    entries=[{'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        for p in sorted(OUT.iterdir()) if p.is_file() and p.name not in ['MANIFEST.json','URP_Figures_Part1_20261020.zip']]
    (OUT/'MANIFEST.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2))
    target=OUT/'URP_Figures_Part1_20261020.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for p in sorted(OUT.iterdir()):
            if p.is_file() and p!=target:z.write(p,'URP_Figures_Part1_20261020/'+p.name)
    with zipfile.ZipFile(target) as z:
        assert z.testzip() is None
        for e in entries:assert hashlib.sha256(z.read('URP_Figures_Part1_20261020/'+e['file'])).hexdigest()==e['sha256']
    print('Brief and verified archive built:',target.stat().st_size,'bytes')


if __name__=='__main__':build()
