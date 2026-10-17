from pathlib import Path
import re,subprocess
from xml.sax.saxutils import escape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image as RLImage,Table,TableStyle
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from PIL import Image
OUT=Path(__file__).resolve().parent
def pdf_report(text,path):
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='ZHBody',fontName='STSong-Light',fontSize=10,leading=16,spaceAfter=6,wordWrap='CJK'))
    styles.add(ParagraphStyle(name='ZHTitle',fontName='STSong-Light',fontSize=19,leading=27,spaceAfter=15,textColor=colors.HexColor('#25384A'),wordWrap='CJK'))
    styles.add(ParagraphStyle(name='ZH1',fontName='STSong-Light',fontSize=15,leading=23,spaceBefore=15,spaceAfter=9,textColor=colors.HexColor('#176B87'),wordWrap='CJK',keepWithNext=True))
    styles.add(ParagraphStyle(name='ZH2',fontName='STSong-Light',fontSize=12,leading=19,spaceBefore=10,spaceAfter=7,textColor=colors.HexColor('#25384A'),wordWrap='CJK',keepWithNext=True))
    styles.add(ParagraphStyle(name='ZHSmall',fontName='STSong-Light',fontSize=7.5,leading=11,wordWrap='CJK'))
    def clean(s):
        s=re.sub(r'\*\*|`','',s).replace('R²','R^2').replace('10⁻¹²','10^-12').replace('m³','m^3').replace('−','-')
        s=s.translate(str.maketrans({'²':'2','³':'3','¹':'1','⁻':'-'}))
        return escape(s)
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=36,leftMargin=36,topMargin=40,bottomMargin=38,
        title=path.stem,author='URP 数据复核与论文框架')
    body=[];lines=text.splitlines();i=0;width=A4[0]-72
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        im=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if im:
            pp=OUT/im.group(2);w,h=Image.open(pp).size;h2=width*h/w
            body.append(RLImage(str(pp),width=width,height=h2));body.append(Spacer(1,5));continue
        if line.startswith('|'):
            tablelines=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):tablelines.append(lines[i].strip());i+=1
            data=[]
            for rr in tablelines:
                cells=[x.strip() for x in rr.strip('|').split('|')]
                if all(re.fullmatch('[: -]+',c) for c in cells):continue
                data.append([Paragraph(clean(c),styles['ZHSmall']) for c in cells])
            if data:
                cols=len(data[0]);t=Table(data,colWidths=[width/cols]*cols,repeatRows=1,hAlign='LEFT')
                t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDE8EE')),('VALIGN',(0,0),(-1,-1),'TOP'),
                    ('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#176B87')),('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#AABBC8')),
                    ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F4F7F9')]),('LEFTPADDING',(0,0),(-1,-1),4),
                    ('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
                body.extend([t,Spacer(1,10)])
            continue
        level=len(line)-len(line.lstrip('#'))
        if level:style=styles['ZHTitle' if level==1 else ('ZH1' if level==2 else 'ZH2')];line=line.lstrip('# ').strip()
        else:style=styles['ZHBody']
        if line.startswith('- '):line='• '+line[2:]
        body.append(Paragraph(clean(line),style))
    def footer(canvas,doc):
        canvas.setFont('STSong-Light',8);canvas.setFillColor(colors.HexColor('#627383'))
        canvas.drawString(36,22,'URP  |  证据复核与论文框架  |  2026-10-17')
        canvas.drawRightString(A4[0]-36,22,str(doc.page))
    doc.build(body,onFirstPage=footer,onLaterPages=footer)

def word_report(mdname):
    from docx import Document
    from docx.shared import Pt,Cm,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    # Pandoc supplies semantic headings, editable tables, and embedded figures.
    target=OUT/(Path(mdname).stem+'.docx')
    subprocess.run(['pandoc',str(OUT/mdname),'-o',str(target),'--resource-path',str(OUT)],check=True,capture_output=True)
    d=Document(target)
    for sec in d.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
        p=sec.footer.paragraphs[0];p.alignment=2
        rr=p.add_run('URP  |  ');rr.font.size=Pt(8)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    for style in d.styles:
        if style.type==1:
            style.font.name='Calibri';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    for name in ['Normal','Body Text','First Paragraph']:
        if name in d.styles:
            d.styles[name].font.size=Pt(10.5);d.styles[name].paragraph_format.line_spacing=1.25
    for name in ['Heading 1','Heading 2','Heading 3']:
        if name in d.styles:d.styles[name].font.color.rgb=RGBColor.from_string('176B87')
    for t in d.tables:
        if 'Table Grid' in d.styles:
            t.style='Table Grid'
        else:
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                border=OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','C6D1D9')]:
                    border.set(qn('w:'+key),value)
                borders.append(border)
            t._tbl.tblPr.append(borders)
        for cell in t.rows[0].cells:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DDE8EE');cell._tc.get_or_add_tcPr().append(sh)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:r.font.size=Pt(8)
    # Keep images within the A4 text width even if the PNG contains wide panels.
    maxwidth=Cm(17.0)
    for sh in d.inline_shapes:
        if sh.width>maxwidth:
            ratio=maxwidth/sh.width;sh.height=int(sh.height*ratio);sh.width=maxwidth
    d.save(target)
