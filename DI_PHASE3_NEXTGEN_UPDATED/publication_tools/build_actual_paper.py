"""Chinese Word manuscript framework filled with actual Sobol/NSGA results."""
from pathlib import Path
import json
import math
import subprocess
import sys
import os
from copy import deepcopy
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(PROJECT))
from model_adapter import load_study,read
DEFAULT_OUT=PROJECT/'actual_publication_20261019'
if not DEFAULT_OUT.exists() and (PROJECT.parent/'publication').exists():DEFAULT_OUT=PROJECT.parent/'publication'
OUT=Path(os.environ.get('URP_PUBLICATION_DIR',str(DEFAULT_OUT)));T=OUT/'tables'
RUN=PROJECT/'output_real_20261019'
if not RUN.exists() and (PROJECT.parent/'complete_runs').exists():RUN=PROJECT.parent/'complete_runs'
FIG=OUT/'figures'
if not FIG.exists():FIG=OUT.parent/'figures'
YEARS=[2019,2020,2021,2022]
SERIF='Noto Serif CJK SC';SANS='Noto Sans CJK SC'
D=Document();TEXT=[]
FORMULAS={
    'parameters':r'\boldsymbol{\theta}=(CC_x,CGC,CDC,S_{sen},K_{cb},WP_e,WP_l,HI_0)',
    'phase1':r'p_1(\boldsymbol{\theta}_1\mid D_1)\propto L_1(D_1\mid\boldsymbol{\theta}_1)p_1(\boldsymbol{\theta}_1)',
    'phase2':r'p_2(\boldsymbol{\theta}_2\mid D_2,\boldsymbol{\theta}_1^*)\propto L_2(D_2\mid\boldsymbol{\theta}_2,\boldsymbol{\theta}_1^*)p_2(\boldsymbol{\theta}_2)',
    'likelihood':r'\log L=C-\frac{1}{2}\sum_jw_j\left(\frac{z_j^{sim}-z_j^{obs}}{\sigma_j}\right)^2',
    'production':r'ET_t=Tr_t+Es_t,\qquad B_H\propto\sum_{t=1}^{H}WP_{eff,t}\frac{Tr_t}{ET_{0,t}}',
    'wp_yield':r'\begin{aligned}WP_{eff}(t)&=WP_e\ (t<t_s),\quad WP_{eff}(t)=WP_l\ (t\geq t_s)\\Y_{13}&=\frac{HI_{adj,H}B_H}{0.87}\end{aligned}',
    'water':r'I+P+CR+GW=Tr+Es+RO+DP+TD+\Delta S+\Delta S_{surf}',
    'sobol':r'S_i=\frac{\mathrm{Var}_{\theta_i}(\mathrm{E}(Y\mid\theta_i))}{\mathrm{Var}(Y)},\qquad S_{T_i}=\frac{\mathrm{E}_{\theta_{-i}}(\mathrm{Var}(Y\mid\theta_{-i}))}{\mathrm{Var}(Y)}',
    'optimization':r'\begin{aligned}\min_{\mathbf a}\;&\big(-Y_{13}(\boldsymbol\theta^*,\mathbf a),I_{act}(\mathbf a)\big)\\\mathrm{s.t.}\;&I_{act}\leq I_{DI},\quad Y_{13}\geq0.95Y_{obs}\end{aligned}',
    'storage':r'\mathcal E_{sameS}=\{r\in\mathcal E:Y_{13,r}\geq0.95Y_{obs},\ I_r\leq I_{DI},\ S_{end,r}\geq S_{end,DI}\}'
}
MATH_BLOCKS={}


def prepare_equations():
    tmp=OUT/'equations_formatting_resource.docx'
    source='\n\n'.join('$$\n'+v+'\n$$' for v in FORMULAS.values())
    subprocess.run(['pandoc','-f','markdown+tex_math_dollars','-o',str(tmp)],input=source,text=True,capture_output=True,check=True)
    snippets=Document(tmp)
    blocks=[deepcopy(p._p) for p in snippets.paragraphs if p._p.xpath('.//m:oMath')]
    if len(blocks)!=len(FORMULAS):raise RuntimeError('Editable equation conversion incomplete')
    MATH_BLOCKS.update(zip(FORMULAS,blocks));tmp.unlink()


def equation(key):
    from docx.text.paragraph import Paragraph
    block=deepcopy(MATH_BLOCKS[key]);D._body._body.insert_element_before(block,'w:sectPr')
    p=Paragraph(block,D._body);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.keep_together=True
    TEXT.append('$$\n'+FORMULAS[key]+'\n$$')


def font(run,size=11,bold=False,color=None):
    run.font.name=SERIF;run.font.size=Pt(size);run.bold=bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),SERIF)
    if color:run.font.color.rgb=RGBColor.from_string(color)


def paragraph(text='',bold=False,size=11):
    p=D.add_paragraph();font(p.add_run(text),size,bold)
    p.paragraph_format.line_spacing=Pt(max(16.5,size*1.45))
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.widow_control=True
    TEXT.append(text);return p


def heading(text,level=1):
    p=D.add_heading(text,level=level)
    for r in p.runs:font(r,15 if level==1 else 12.5,True,'24556B')
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.line_spacing=Pt(22 if level==1 else 18)
    p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(7)
    TEXT.append('\n'+text);return p


def table(title,headers,rows):
    paragraph(title,True,10.5)
    t=D.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for c,h in zip(t.rows[0].cells,headers):c.text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    for i,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=Pt(14)
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:font(r,10.5,i==0)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EFF3');c._tc.get_or_add_tcPr().append(shade)
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    TEXT.append('\n'+title+'\n'+' | '.join(headers)+'\n'+'\n'.join(' | '.join(map(str,r)) for r in rows))
    D.add_paragraph();return t


def dimensions(section,wide=False):
    section.orientation=WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
    section.page_width=Cm(29.7 if wide else 21);section.page_height=Cm(21 if wide else 29.7)
    section.left_margin=section.right_margin=Cm(2)
    section.top_margin=section.bottom_margin=Cm(1.7)
    section.header_distance=section.footer_distance=Cm(.8)
    grid=section._sectPr.find(qn('w:docGrid'))
    if grid is not None:section._sectPr.remove(grid)
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if not foot.text:
        font(foot.add_run('URP · 四年真实模型结果   '),9)
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)


def main():
    OUT.mkdir(exist_ok=True);study=load_study()
    status=read(RUN/'EXPERIMENT_STATUS.json')
    if status['status']!='COMPLETE':raise RuntimeError('Experiment incomplete')
    prepare_equations()
    registry=read(OUT/'FIGURE_REGISTRY.json')['figures']
    chosen=pd.read_csv(T/'NSGA_main_candidates.csv').set_index('Year')
    sensitivity=pd.read_csv(T/'Sobol_all_seed_indices.csv')
    stability=pd.read_csv(T/'Sobol_numerical_stability.csv')
    cond=pd.read_csv(T/'NSGA_conditional_validation.csv')
    hist=pd.read_csv(T/'NSGA_convergence_history.csv')
    saving=pd.read_csv(T/'Irrigation_saving_decomposition.csv').set_index('Year')
    schedules=pd.read_csv(T/'NSGA_event_schedules.csv')
    dimensions(D.sections[0]);D.core_properties.title='URP四年Sobol与NSGA-III论文写作框架（已填真实结果）'
    D.core_properties.subject='2019—2022滴灌玉米；冻结MCMC；真实年度AquaCrop响应'
    D.core_properties.author='URP项目'
    for style in ['Normal','Body Text']:
        D.styles[style].font.name=SERIF;D.styles[style].font.size=Pt(11)
    heading('四年滴灌玉米敏感性分析与灌溉优化',1)
    paragraph('URP论文写作框架与真实结果正文稿',True,17)
    paragraph('项目：基于马尔科夫链的不同灌溉方式灌溉制度优化\n项目编号：X2026100190937\n数据年份：2019、2020、2021、2022\n本稿用于组织论文与解释图表；研究证据集中于DI，题目和结论不扩展到未经比较的灌溉方式。')
    paragraph('本次新增的Sobol指数、逐次模型响应、NSGA-III/NSGA-II前沿、每代种群、候选回放及条件扰动均来自本轮原年度模型计算。原MCMC样本保持冻结。既有MCMC严格收敛诊断限制仍保留，完成下游计算不等于原链已收敛。',True)
    heading('建议题目与论文主线',2)
    paragraph('建议题目：分阶段贝叶斯校准与敏感性分析支持的滴灌玉米灌溉制度优化——基于2019—2022年资料。论证顺序是过程依据、参数贡献、决策权衡、适用限制。参考Liu等（2025）的MCMC—AquaCrop—多目标优化组织方式，但不移用棉田参数、观测或节水幅度。')
    paragraph('如果论文重点放在本轮发现的季末储水影响，也可用“季末土壤储水条件影响下的滴灌玉米灌水量—产量权衡”作为题目方向。创新应体现在识别灌水节省的水量来源、端点条件与参数不确定性，而不是仅组合三个算法名称。')
    heading('摘要示例',2)
    values='；'.join(f'{y}年{chosen.loc[y,"Water_saving_pct"]:.1f}%' for y in YEARS)
    paragraph('为研究DI玉米参数不确定性与灌水分配的联系，以2019—2022年气象、冠层、生物量和产量资料及既有分阶段MCMC成果为基础，复现四年改进AquaCrop模型基准，在公共独立均匀范围内开展八参数Sobol分析，并固定年度参数中心点进行双目标灌溉优化。敏感性分析由N=128、三个scramble筛查扩展至N=512；优化采用100个体、100代、五随机种子的NSGA-III，并与同预算NSGA-II比较。在产量不低于年度发表观测95%、实际灌水不超过原DI的约束下，最低灌水量候选的模拟节水率依次为'+values+'。候选另作两个参数组的整行条件扰动，未将其视为八维联合后验。结果用于说明锁定初始水分与固定灌水日期条件下的模拟权衡；灌水减少量还须分解为耗水、排水及季末储水变化，不能直接解释为长期可持续节水。')
    paragraph('关键词：滴灌玉米；AquaCrop；分阶段MCMC；Sobol；NSGA-III；水量平衡')
    heading('1 引言')
    paragraph('第一段解释研究问题：灌溉水量既影响根区供水与作物胁迫，也影响季末储水和非生产性蒸发。仅以年灌水量或单一产量评价管理，难以说明节水发生在哪个过程。背景中的水资源比例、研究区气候与试验信息，应引用确切来源，不填入未经核对的宏观数字。')
    paragraph('第二段解释模型校准的必要性：冠层增长、作物蒸腾、地上生物量形成和收获指数相互联系，多个参数可能补偿同一观测误差。因此，本研究使用阶段一冠层—水分参数及阶段二生物量—产量参数的条件校准成果，并同时报告其采样局限。')
    paragraph('第三段给出研究缺口与问题：在声明的独立范围下，参数对不同输出的方差贡献是否随年份改变？在观测产量门槛下，灌水量与产量有哪些可行权衡？候选的条件参数扰动和季末储水代价如何？这些问题比简单列出三种算法名称更能构成论文主线。')
    heading('2 材料与方法')
    table('表1 四年基准及绝对产量门槛',['年份','发表Y13','模型基准Y13','原DI / mm','95%观测门槛'],[
        [y,f'{study["observed_Y13"][str(y)]:.2f}',f'{study["baseline_reference"][str(y)]["Y_13pct"]:.4f}',
         f'{study["baseline_irrigation_mm"][str(y)]:.2f}',f'{.95*study["observed_Y13"][str(y)]:.4f}'] for y in YEARS])
    calendar=pd.read_csv(T/'Annual_model_calendar.csv')
    table('年度日历与灌溉事件配置',['年份','播种日','日历截止界','计算日数','WP切换日','灌水事件数'],[
        [int(r.Year),r.Planting_date,r.Calendar_cutoff,int(r.Calendar_days),r.WP_switch,int(r.Irrigation_events)] for r in calendar.itertuples()])
    paragraph('Y13单位为t/ha，含水率13%；干基产量Ydry=0.87Y13。Field LAI转换CC和数字化生物量参与校准，不能再次称为独立验证资料。遥感时序只作辅助比较。研究地点、品种、田间重复、数字化误差和LAI转换来源，仍应从原始试验材料补齐。四年本轮实际使用的气象及B组ETo输入均来自新补年度包，未再以年度比例缩放旧驱动代替。')
    heading('2.1 既有MCMC与模型结构',2)
    paragraph('阶段一参数为CCx、CGC、CDC、SenescenceCD、Kcb；阶段二参数为WP early、WP late、HI0。阶段二条件于阶段一交接点，不构成八参数联合后验。本轮使用已核验保留样本的年度中心参数作确定性决策输入，不重新采样MCMC。相关保留行数不能当作独立有效样本量；原严格Rhat/ESS和链长检查未全部通过的事实应进入补充材料和讨论。')
    paragraph('均值参数点的确定性模拟一般不等于后验预测均值；文中模型基准是前者，不将其当作完整联合后验的平均预测。')
    equation('parameters');equation('phase1');equation('phase2');equation('likelihood')
    paragraph('这里给出目标分布与加权高斯工作误差的解释形式，实际权重、误差尺度及交接点以原采样合同为准。Tr目标若由ET−Es构成，其与Es误差不必独立；ET不应再作为额外独立似然重复计权。后续可显式建模相关误差或使用ET/Es联合观测。本轮未修改原似然或重跑链。')
    paragraph('作物过程仍使用原年度Modified DI AquaCrop 3.1.0。早、晚有效WP按各年冻结切换日设置，干生物量与最终调整收获指数形成干基籽粒产量。WP的两段结构属于本项目的模型改动；WP early与WP late的差异也可能吸收驱动或结构误差，不能只作生理解释。')
    equation('production');equation('wp_yield')
    paragraph('H为实际收获步，ts为固定年度切换日；WP—Tr/ETo关系是过程解释式，不替代原代码中的胁迫与CO₂处理。HIadj为实际调整后的收获指数，不能直接用输入HI0替代。SenescenceCD遵循原引擎取最近整数日。')
    paragraph('土壤计算使用24×0.1 m网格。2019—2021保留原保守映射，2022按冻结合同保留旧插值初始化以复现原成果。初始含水配置与剖面扩展不能直接当作全深度独立实测；论文应列出观测深度、深层水力参数及外推方法。年际比较也受这类初始状态和参数化差异影响。')
    paragraph('为使公共参数范围内提前成熟或死亡的组合有明确响应，记录原模型实际收获产量与生物量，继续土壤水量平衡至各年原日历截止日；收获后按AquaCrop休耕逻辑停止作物蒸腾及后续灌水，保留声明的覆膜管理。ET因此为统一日历窗口内Tr+Es；同时导出实际收获日期、在田天数、提前结束标志及收获前Tr/Es。旧接口失败设计未删行或填零。')
    paragraph('数值加速仅编译原数值方程并复用首次可能灌水前的相同状态，不使用代理模型或近似插值。四年中心均复现原Y13、Tr、Es、ET基准；28个中心、全范围参数及管理对照案例的原方程与加速结果最大绝对差约1.14×10⁻¹³。每次响应都检查土壤物理界限与完整逐日水量账。')
    equation('water')
    paragraph('I为实际灌水，P为降水，CR/GW为毛管/地下水入流，RO/DP/TD为径流、深渗和暗管排水；ΔS与ΔSsurf分别为期末减期初的土壤和地表储水变化。完整通量逐项保留；结构性零通量按源模型设置报告。')
    heading('2.2 Sobol输入、输出与稳定性',2)
    equation('sobol')
    paragraph('S1描述单个独立输入的主效应，ST还包括与其他输入的交互贡献。指数相对于给定输入分布下的输出方差归一化，不是物理因果回归系数，也不能跨输出单位直接解释绝对影响量。')
    table('表2 公共独立范围（不是原MCMC先验密度）',['参数','下界','上界'],[[k,*study['global_physical_bounds'][k]] for k in study['global_physical_bounds']])
    paragraph('采用SALib的scrambled Sobol/Saltelli设计，不估计二阶指数；八参数时每年每种子10N行。先运行N=128、种子101/202/303，共15,360行，随后四年均扩展到N=512。扩展设计与原设计前缀逐位相同，精确复用15,360行已计算响应，新增46,080次响应，最终61,440个设计行全部有真实响应。响应包括Y13、B、Tr、Es、ET、MaxCC；保存S1、ST、512次bootstrap的95%区间半宽及输出方差。')
    paragraph('主分析比较ST，并保留三个种子的独立估计。筛查门槛为前三参数集合一致且最大ST区间半宽≤0.1；N=512后仍未通过的项单独列出，不以平均热图掩盖，也不强制排名。S1的微小负估计及ST−S1的不等式偏差属于有限样本误差的一部分，不能随意截断或据此宣称校准吸收了某比例的信息。')
    heading('2.3 NSGA-III决策变量与公平对照',2)
    equation('optimization')
    paragraph('目标为min(−Y13,Iactual)，约束为Iactual≤年度原DI灌水量、Y13≥0.95×年度发表观测。四个有序灌溉窗口由原灌水事件按时间顺序分组，窗口内保持原事件比例和日期，只优化各组总量。该控制不是SMT，也不是优化灌期或凭空增加灌水次数。接口保留计划量、实际量及提前收获后取消量。')
    paragraph('每算法100个体、100代，随机种子31/73/109/151/199，年度分别寻优，NSGA-III使用100个双目标参考方向；同预算NSGA-II作为对照。初始化含真实基准及随机预算内分配，修复算子把超预算计划总量投回年度水量单纯形；两算法使用同样机制。每代保存X/F/G/CV，每次评估保存真实输出。HV按Y13/观测、I/原DI归一化，参考点[0,1.01]，仅作同年度算法比较。')
    paragraph('从NSGA-III五种子非支配并集中选ε门槛内最低灌水量主候选，并列出两个相邻可行备选；该选取不是几何膝点。候选独立回代核对目标与约束。对阶段一、阶段二分别整行抽取128组冻结参数，固定另一组中心，保留组内相关；不把两组拼成联合后验。抽样索引独立有放回，二项区间仅描述给定经验样本池下的Monte Carlo误差，不包含原MCMC误差或田间风险。')
    heading('3 结果与图表解读')
    heading('3.1 参数贡献及年际差异',2)
    parameter_names={'WP_late':'晚期WP','WP_early':'早期WP','HI0':'HI0','Kcb':'Kcb','CCx':'CCx',
        'CGC_CD':'CGC','CDC_CD':'CDC','SenescenceCD':'衰老起始日'}
    for output,label in [('Y_13pct','产量Y13'),('ET_mm','总耗水ET'),('B_t_ha','地上干生物量B'),('Tr_mm','蒸腾Tr'),('Es_mm','土壤蒸发Es'),('MaxCC','最大冠层覆盖度')]:
        rows=[]
        for year in YEARS:
            d=sensitivity[(sensitivity.Year==year)&(sensitivity.Output==output)].groupby('Parameter').ST.mean().nlargest(3)
            rows.append([year]+[f'{parameter_names[k]} ({v:.3f})' for k,v in d.items()])
        table(f'{label}的前三ST均值（原始区间见CSV）',['年份','第一','第二','第三'],rows)
    paragraph('写作时先报告ST较大且跨种子较稳定的参数，再解释它们与输出过程的联系。WP和HI0对产量的贡献分别联系生物量形成与籽粒分配；Kcb联系蒸腾及总耗水。某参数对一个输出贡献小，并不等于它在其他输出或另一输入分布下没有作用。不同年份驱动与基准水分不同，是年际贡献变化的背景；不能只凭相关或名次差异作因果归因。')
    failed=stability[~stability.Stable_for_screening]
    if len(failed):
        paragraph(f'N=512后仍有{len(failed)}个“年份×输出”未同时满足预设筛查门槛，详见下表。论文对这些项保留不确定性说明；尤其相邻参数名次可能在区间重叠时互换，不能写为已证明稳定排序。',True)
        table('仍需保留数值不确定性的输出',['年份','输出','前三集合一致','最大ST半宽'],[
            [int(r.Year),r.Output,'是' if r.Same_top_set_across_scrambles else '否',f'{r.Max_ST_95ci_halfwidth:.3f}'] for r in failed.itertuples()])
    else:paragraph('N=512后，全部24个“年份×输出”均通过预设前三集合和区间半宽筛查。该检查支持本次范围下的数值稳定性，不等于输入范围或模型结构已经获得独立验证。')
    heading('3.2 真实Pareto前沿与代表方案',2)
    table('表3 ε门槛内最低灌水量主候选',['年份','I候选 mm','Y13候选','节水率 %','相对观测产量 %','IWP kg/m³'],[
        [y,f'{chosen.loc[y,"Irrigation_mm"]:.2f}',f'{chosen.loc[y,"Y13"]:.4f}',f'{chosen.loc[y,"Water_saving_pct"]:.2f}',
         f'{chosen.loc[y,"Yield_change_vs_observed_pct"]:+.2f}',f'{chosen.loc[y,"IWP_kg_m3"]:.3f}'] for y in YEARS])
    for year in YEARS:
        c=chosen.loc[year]
        paragraph(f'{year}年主候选灌水量为{c.Irrigation_mm:.2f} mm、模拟产量为{c.Y13:.4f} t/ha，相对原DI节水{c.Water_saving_pct:.2f}%，相对发表观测产量变化{c.Yield_change_vs_observed_pct:+.2f}%。该方案是在允许产量降至观测95%的ε约束下选得，不能写成“节水且产量完全不变”。对应窗口量、逐次实际灌水和两个备选均在结果表中。')
    rows=[]
    for year in YEARS:
        for algorithm in ['NSGA3','NSGA2']:
            h=hist[(hist.Year==year)&(hist.Algorithm==algorithm)]
            end=h.sort_values('Generation').groupby('Seed').tail(1)
            rows.append([year,algorithm,len(end),int(h.Evaluations.max()),f'{end.Hypervolume.median():.5f}',
                f'{end.Hypervolume.min():.5f}–{end.Hypervolume.max():.5f}'])
    table('表4 同预算算法收敛记录',['年份','算法','种子数','每种子评估数','末代HV中位数','五种子HV范围'],rows)
    paragraph('报告两算法的可行率、前沿分布与末代HV范围；若NSGA-II在双目标任务中相当或更好，应如实描述。NSGA-III的名称本身不构成优越性证据。这里的五种子重复检验算法的随机性，未构成对未来气候或其他田块的外部验证。')
    paragraph('本轮依据补齐年度源码和原日历，使用固定四窗口控制及观测95%门槛；与旧图的模型、控制或风险口径不同。新旧节水率的数值差异不能归因为算法性能提升。收获后土壤储水降低也未结合休耕期补给核算，单季灌水减少不能直接推断多年总耗水节约。')
    heading('3.3 条件参数扰动与尾部产量',2)
    c1=cond[cond.Candidate_ID.str.endswith('_C1')]
    table('表5 主候选的分阶段条件扰动',['年份','扰动参数组','均值Y13','Q05 Y13','门槛达标比例','重采样MCSE'],[
        [int(r.Year),'阶段一' if r.Group.startswith('phase1') else '阶段二',f'{r.Mean_Y13:.3f}',f'{r.Q05_Y13:.3f}',
         f'{r.Probability_above_absolute_gate:.3f}',f'{r.MC_standard_error_conditional:.3f}'] for r in c1.itertuples()])
    paragraph('主候选在确定性中心点满足绝对产量门槛，不意味着在条件参数扰动下以高概率满足门槛。若尾部产量或条件达标比例偏低，应把该方案描述为“中心点节水极限候选”，并比较其相邻备选，或在后续研究中预先规定概率约束重新优化。不能在获得当前结果后把它追认为已做过稳健优化。')
    heading('3.4 节水来源与季末储水',2)
    table('表6 灌水减少的水量分解（mm）',['年份','灌水减少','耗水减少','季末土壤储水减少','排水/径流减少'],[
        [y,f'{saving.loc[y,"Irrigation_saving_mm"]:.2f}',f'{saving.loc[y,"ET_reduction_mm"]:.2f}',
         f'{saving.loc[y,"Reduced_soil_storage_gain_mm"]:.2f}',f'{saving.loc[y,"Loss_reduction_mm"]:.2f}'] for y in YEARS])
    paragraph('完整分解另含地表储水及其他入水差额，全部项之和与灌水减少量一致。季末土壤储水减少项为基准储水增量减去候选储水增量；正值说明候选减少了季末储水，或增加了土壤水消耗。因此，灌水节省率与作物耗水减少率必须分别报告。当前优化没有季末储水下限，不能据此宣称多年连作条件下同样节水。')
    storage=pd.read_csv(T/'Terminal_storage_screening.csv')
    screen_rows=[]
    for y in YEARS:
        r=storage[(storage.Year==y)&(storage.Screen=='same_baseline_end_storage')].iloc[0]
        screen_rows.append([y,f'{r.Storage_floor_mm:.2f}',f'{r.Irrigation_mm:.2f}',f'{r.Y13:.4f}',f'{r.Water_saving_pct:.2f}',f'{r.Algorithm}/{int(r.Seed)}'])
    table('表7 追加季末储水不低于基准条件后的已评估点筛查',['年份','储水门槛mm','筛查I mm','Y13','节水%','来源算法/种子'],screen_rows)
    equation('storage')
    paragraph('表7在原40次搜索的所有已评估可行点中追加储水条件，再取最低灌水量并回代。它提供更接近相同季末水分条件的对照，但不等于带储水约束重新优化，也不证明该约束下的最优值。另一项“不低于初始储水”筛查及未找到点的情况见完整表，不能把未找到解释为数学上不可行。下一轮应先根据田间需求确定储水门槛，再把它写入约束重新寻优。')
    max_event=schedules.groupby('Year').Candidate_actual_mm.max()
    paragraph('候选的最大单次灌水量依次为'+ '、'.join(f'{y}年{max_event.loc[y]:.1f} mm' for y in YEARS)+'。本轮没有另行假定设备的单次应用上限；实际推广前需依据田间设备、入渗能力和轮灌限制补入约束，并重新求解。')
    heading('4 讨论、完善方向与论文边界')
    paragraph('模型层面：补齐试验地点、品种、重复数和误差来源，解释两阶段有效WP与覆膜处理的农艺依据；使用真正独立的年份或处理检验过程预测。气象、ETo修正、LAI转换和数字化生物量误差可能在参数之间补偿，不能把高拟合优度等同机制已被唯一识别。')
    paragraph('推断层面：保留原MCMC严格诊断，后续可通过更长独立ensemble、可辨识性分析或联合建模改善。若要传播八参数联合不确定性，需获得相应联合样本或显式建立分层模型；当前两个条件组分别扰动的结果只能支持各自条件下的敏感程度说明。')
    paragraph('决策层面：下一轮可加入季末土壤储水、单次灌水量和设备能力约束，以比较单季极限节水与可持续管理。若研究通用制度，需留一年或嵌套跨年验证，避免在相同四年既寻优又宣称独立验证。若增加天气不确定性、CVaR或概率约束，应先定义风险目标并重新优化，不能把当前确定性前沿加上误差棒后改称联合稳健最优。')
    paragraph('Sobol层面：公共独立范围支持参数方差分解，但不代表MCMC先验密度或相关后验。需要研究相关后验贡献时，可考虑适用于依赖输入的Shapley效应等方法；要比较先验和后验不确定性，应同时报告输出方差变化并说明分布及依赖结构，不能只将两组归一化ST的高低解释为信息吸收。')
    heading('5 结论写作示例')
    paragraph('四个年度原模型在冻结中心参数下复现既有产量与水分基准，并在统一日历水量核算下完成全范围真实Sobol分析。参数贡献随输出和年份而变化，应结合数值区间解读。固定原灌水日期、允许观测产量最多5%降幅时，四年均获得可行的灌水—产量权衡方案；主候选模拟节水率为'+values+'。这些幅度仅适用于锁定初始水分和当前约束，条件参数扰动、季末储水代价及无独立田间验证的限制必须与节水数字共同报告。')
    heading('参考文献与软件依据')
    for ref in [
        'Liu Y, Wang Y, Šimůnek J, Liao R. 2025. An integrated data assimilation, crop modeling, and multi-objective optimization framework for improving cotton irrigation water use efficiency. Agricultural Water Management, 319:109774. DOI:10.1016/j.agwat.2025.109774.',
        'FAO. AquaCrop模型参考手册；本轮实现为AquaCrop Python 3.1.0及随年度包提供的Modified DI源代码。',
        'Saltelli A, et al. 2010. Variance based sensitivity analysis of model output: Design and estimator for the total sensitivity index. Computer Physics Communications, 181:259–270. DOI:10.1016/j.cpc.2009.09.018.',
        'Deb K, Jain H. 2014. An evolutionary many-objective optimization algorithm using reference-point-based nondominated sorting approach, Part I. IEEE Transactions on Evolutionary Computation, 18:577–601. DOI:10.1109/TEVC.2013.2281535.',
        'Foreman-Mackey D, et al. 2013. emcee: The MCMC Hammer. PASP, 125:306–312. DOI:10.1086/670067.',
        'Herman J, Usher W. 2017. SALib: An open-source Python library for sensitivity analysis. JOSS. DOI:10.21105/joss.00097.',
        'Blank J, Deb K. 2020. pymoo: Multi-Objective Optimization in Python. IEEE Access. DOI:10.1109/ACCESS.2020.2990567.'
    ]:paragraph(ref,size=10.5)
    heading('附录：全部新增图件及逐图说明')
    paragraph('PNG用于Word排版，PDF用于矢量输出。各图覆盖四年；横向分面图单独使用横向页面，保持文字大小。原始数值位于tables，全部种群与逐次响应位于运行结果包。')
    for number,entry in enumerate(registry,1):
        wide=entry['orientation']=='landscape'
        section=D.add_section(WD_SECTION_START.NEW_PAGE);dimensions(section,wide)
        heading(f'图{number}  {entry["title"]}',2)
        p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIG/(entry['name']+'.png')),width=Cm(24.8 if wide else 16.7))
        paragraph(entry['caption'],size=10.5)
        paragraph('源图：'+entry['name']+'；完整数值与稳定性记录随包提供。',size=10.5)
    atlas=read(OUT/'BASELINE_FIGURE_ATLAS.json')
    for entry in atlas:
        old_id,title,name,orientation,caption,*_=entry
        if name in ['F01_workflow','F10_climate']:continue
        dest=OUT/'baseline_word_figures'/(name+'.png')
        section=D.add_section(WD_SECTION_START.NEW_PAGE);dimensions(section,orientation=='wide')
        heading('既有校准图 · '+title,2)
        p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(dest),width=Cm(24.8 if orientation=='wide' else 16.7))
        paragraph(caption,size=10.5)
        paragraph('原MCMC及过程数据保持冻结；本图提供校准背景，不作为新Sobol或优化结果。',size=10.5)
    target=OUT/'URP_Sobol_NSGAIII_paper_with_actual_results.docx';D.save(target)
    (OUT/'URP_Sobol_NSGAIII_paper_with_actual_results.md').write_text('\n\n'.join(TEXT),encoding='utf-8')
    print(target,flush=True)


if __name__=='__main__':main()
