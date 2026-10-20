"""Readable Excel summary, with full-precision source CSV/NPZ kept separately."""
from pathlib import Path
import os
import pandas as pd
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

PROJECT=Path(__file__).resolve().parents[1]
DEFAULT_OUT=PROJECT/'actual_publication_20261019'
if not DEFAULT_OUT.exists() and (PROJECT.parent/'publication').exists():DEFAULT_OUT=PROJECT.parent/'publication'
OUT=Path(os.environ.get('URP_PUBLICATION_DIR',str(DEFAULT_OUT)));T=OUT/'tables'
SHEETS={
    '年度模型日历':'Annual_model_calendar.csv',
    'Sobol_完整指数':'Sobol_all_seed_indices.csv',
    'Sobol_四年汇总':'Sobol_four_year_summary.csv',
    'Sobol_稳定性':'Sobol_numerical_stability.csv',
    'NSGA_主候选':'NSGA_main_candidates.csv',
    'NSGA_三个候选':'NSGA_epsilon_candidates.csv',
    'NSGA_全部前沿':'NSGA_all_seed_Pareto_points.csv',
    'NSGA_收敛记录':'NSGA_convergence_history.csv',
    'NSGA_算法对照':'NSGA_algorithm_comparison.csv',
    'NSGA_逐次灌水':'NSGA_event_schedules.csv',
    '条件扰动核验':'NSGA_conditional_validation.csv',
    '完整水量账':'Water_ledger_baseline_and_candidate.csv',
    '节水量分解':'Irrigation_saving_decomposition.csv',
    '季末储水筛查':'Terminal_storage_screening.csv'}


def main():
    descriptions=[
        ['研究范围','2019、2020、2021、2022年DI玉米；真实原年度AquaCrop模型。'],
        ['MCMC','原样冻结，不把两组条件样本拼成八维联合后验；严格诊断限制保留。'],
        ['Sobol','8参数公共独立均匀范围；N512、3种子；S1/ST为原始估计，不裁剪负S1。'],
        ['区间','S1_conf/ST_conf是95%bootstrap半宽，不是上下界。'],
        ['稳定性','23/24项通过；2022产量的WP early/Kcb贡献接近，保留排序不确定性。'],
        ['优化','NSGA-III与NSGA-II各100个体×100代×5种子；固定年度中心点。'],
        ['变量','四个有序灌溉窗口总量；原日期及组内水量比例不变，不是灌期优化或SMT。'],
        ['约束','实际I≤年度原DI；Y13≥发表观测的95%。主候选按最低灌水量选，不是几何膝点。'],
        ['产量单位','Y13：13%含水率t/ha；干基产量=0.87Y13。B为地上干生物量t/ha。'],
        ['用水单位','I/Tr/Es/ET为mm；IWP=100Y13/I，单位kg/m³；IWP不同于ET水生产率。'],
        ['条件扰动','每候选每参数组128次；Wilson区间仅为经验样本池重采样MC误差，未覆盖MCMC/田间/模型误差。'],
        ['季末储水','优化未设置季末储水下限，节水量可能包含减少季末储水，不能等同长期耗水节省。'],
        ['精度','Excel显示保留3位小数以便阅读；原CSV与无损NPZ保存计算精度，完整逐次记录见运行结果包。'],
        ['文件结构','本工作簿为阅读汇总；所有61,440行Sobol设计、400,000次优化评估和4,000代种群记录另完整提供。']]
    target=OUT/'URP_Sobol_NSGAIII_results_tables.xlsx'
    with pd.ExcelWriter(target,engine='openpyxl') as writer:
        pd.DataFrame(descriptions,columns=['项目','说明']).to_excel(writer,sheet_name='先读说明',index=False)
        main=pd.read_csv(T/'NSGA_main_candidates.csv').set_index('Year')
        ledger=pd.read_csv(T/'Water_ledger_baseline_and_candidate.csv')
        water=pd.read_csv(T/'Irrigation_saving_decomposition.csv').set_index('Year')
        condition=pd.read_csv(T/'NSGA_conditional_validation.csv')
        screen=pd.read_csv(T/'Terminal_storage_screening.csv')
        primary=[];risk=[]
        for year in [2019,2020,2021,2022]:
            r=main.loc[year];base=ledger[(ledger.Year==year)&(ledger.Scenario=='baseline')].iloc[0]
            primary.append({'年份':year,'原DI灌水_mm':base.Irrigation_mm,'候选灌水_mm':r.Irrigation_mm,
                '发表观测Y13_t_ha':r.Observed_Y13,'候选模拟Y13_t_ha':r.Y13,'灌水减少_百分数':r.Water_saving_pct,
                '产量相对观测变化_百分数':r.Yield_change_vs_observed_pct,'IWP_kg_m3':r.IWP_kg_m3})
            c=condition[(condition.Year==year)&condition.Candidate_ID.str.endswith('_C1')].set_index('Group')
            q=screen[(screen.Year==year)&(screen.Screen=='same_baseline_end_storage')].iloc[0]
            risk.append({'年份':year,'耗水减少_mm':water.loc[year,'ET_reduction_mm'],
                '季末土壤储水减少_mm':water.loc[year,'Reduced_soil_storage_gain_mm'],
                '阶段一条件达标比例':c.loc['phase1_conditional','Probability_above_absolute_gate'],
                '阶段二条件达标比例':c.loc['phase2_conditional','Probability_above_absolute_gate'],
                '相同季末储水筛查点灌水_mm':q.Irrigation_mm,'该筛查点灌水减少_百分数':q.Water_saving_pct,
                '说明':'储水筛查不是带新约束重新寻优；条件比例不是田间稳产概率'})
        pd.DataFrame(primary).to_excel(writer,sheet_name='四年核心结果',index=False)
        pd.DataFrame(risk).to_excel(writer,sheet_name='储水和条件核验',index=False)
        for name,filename in SHEETS.items():pd.read_csv(T/filename,float_precision='round_trip').to_excel(writer,sheet_name=name,index=False)
        book=writer.book
        for ws in book.worksheets:
            ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
            ws.sheet_view.showGridLines=False
            for c in ws[1]:
                c.fill=PatternFill('solid',fgColor='24556B');c.font=Font(name='Microsoft YaHei',size=11,bold=True,color='FFFFFF')
                c.alignment=Alignment(vertical='center',wrap_text=True)
            ws.row_dimensions[1].height=34
            for index,column in enumerate(ws.columns,1):
                name=str(column[0].value)
                ws.column_dimensions[get_column_letter(index)].width=min(29,max(12,len(name)*.8+3))
                for cell in column[1:]:
                    cell.font=Font(name='Microsoft YaHei',size=10)
                    if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F0F5F7')
                    if isinstance(cell.value,float):cell.number_format='0.000'
                    cell.alignment=Alignment(vertical='top',wrap_text=False)
            headers=[c.value for c in ws[1]]
            if 'Stable_for_screening' in headers:
                j=headers.index('Stable_for_screening')+1
                for row in ws.iter_rows(min_row=2):
                    if ws.cell(row[0].row,j).value is False:
                        for cell in row:cell.fill=PatternFill('solid',fgColor='FCE9D8')
        info=book['先读说明'];info.column_dimensions['A'].width=18;info.column_dimensions['B'].width=115
        for row in info.iter_rows(min_row=2):
            row[1].alignment=Alignment(wrap_text=True,vertical='top');info.row_dimensions[row[0].row].height=31
    print(target,flush=True)


if __name__=='__main__':main()
